"""Tests for experiment E1 and the retrieval metrics (no models, no API key)."""

import sys
import tempfile
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from evaluation.e1_retrieval import (  # noqa: E402
    LLM_ONLY_RULES,
    LLMOnly,
    blind_grading_table,
    e1_table,
    manual_scores,
    retrieval_ranks,
    retrieval_table,
)
from evaluation.eval_abstention import run_eval_resumable  # noqa: E402


def doc(source, number, lang="ar"):
    return {"source": source, "article_number": str(number), "article_text": "نص", "language": lang}


def gold_df():
    return pd.DataFrame({
        "qid": ["Q1", "Q2", "Q3", "Q4"],
        "question": ["سؤال 1", "سؤال 2", "Question 3", "U سؤال 4"],
        "language": ["ar", "ar", "en", "ar"],
        "answerable": [True, True, True, False],
        "gold_articles": [["5"], ["24"], ["7"], []],
        "gold_source": ["pdpl", "implementing", "transfer", ""],
    })


def cache():
    return {
        # gold ranked 3rd by BM25, 1st after rerank
        "Q1": {"bm25": [doc("pdpl", 1), doc("pdpl", 2), doc("pdpl", 5)],
               "dense": [doc("pdpl", 5)], "fused": [doc("pdpl", 2), doc("pdpl", 5)],
               "final": [doc("pdpl", 5), doc("pdpl", 2)]},
        # same NUMBER in the wrong document must not count (PDPL 24 is not Implementing 24)
        "Q2": {"bm25": [doc("pdpl", 24)], "dense": [doc("pdpl", 24), doc("implementing", 24)],
               "fused": [doc("implementing", 24)], "final": [doc("implementing", 24)]},
        # BM25 results may carry only source_file
        "Q3": {"bm25": [{"source_file": "transfer_en_articles.csv", "article_number": "7"}],
               "dense": [doc("pdpl", 7, "en")], "fused": [], "final": []},
        "Q4": {"bm25": [], "dense": [], "fused": [], "final": []},
    }


def test_retrieval_ranks_match_document_and_article():
    ranks = retrieval_ranks(cache(), gold_df())
    assert list(ranks["qid"]) == ["Q1", "Q2", "Q3"]            # unanswerable questions excluded
    r = ranks.set_index("qid")
    assert r.loc["Q1", "bm25_rank"] == 3 and r.loc["Q1", "reranked_rank"] == 1
    assert r.loc["Q2", "bm25_rank"] is None or pd.isna(r.loc["Q2", "bm25_rank"])   # wrong document
    assert r.loc["Q2", "dense_rank"] == 2
    assert r.loc["Q3", "bm25_rank"] == 1                        # derived from source_file
    assert pd.isna(r.loc["Q3", "dense_rank"])                   # PDPL 7 is not Transfer 7


def test_retrieval_table_recall_and_mrr_per_language():
    t = retrieval_table(retrieval_ranks(cache(), gold_df()), k=5).set_index(["method", "language"])
    bm25_all = t.loc[("bm25", "all")]
    assert bm25_all["n"] == 3
    assert bm25_all["recall@5"] == pytest.approx(2 / 3, abs=1e-3)      # Q1 and Q3
    assert bm25_all["mrr@5"] == pytest.approx((1 / 3 + 0 + 1) / 3, abs=1e-3)
    assert t.loc[("reranked", "ar"), "recall@5"] == 1.0
    assert t.loc[("reranked", "en"), "recall@5"] == 0.0
    assert pd.isna(t.loc[("reranked", "all"), "recall@10"])            # only 5 kept after rerank


def test_llm_only_sends_no_sources_and_detects_not_covered(monkeypatch):
    sent = {}

    def fake_call_llm(prompt, provider=None, model=None, system=None):
        sent.update(prompt=prompt, system=system)
        return "NOT_COVERED: لا أعرف" if "U سؤال" in prompt else "حسب نظام حماية البيانات (Article 5)"

    monkeypatch.setattr("evaluation.e1_retrieval.call_llm", fake_call_llm)
    out = LLMOnly("ollama").ask("سؤال 1")
    assert sent["system"] == LLM_ONLY_RULES and "Retrieved legal sources" not in sent["prompt"]
    assert not out["abstained"] and out["final_sources"] == [] and out["prompt_style"] == "llm_only"
    assert LLMOnly("ollama").ask("U سؤال 4")["abstained"]


def test_e1_table_compares_the_three_systems(monkeypatch):
    replies = {"سؤال 1": "الإجابة (Article 5)", "سؤال 2": "الإجابة (Article 9)",
               "Question 3": "NOT_COVERED: unknown", "U سؤال 4": "الغرامة 5 ملايين (Article 35)"}
    monkeypatch.setattr("evaluation.e1_retrieval.call_llm",
                        lambda prompt, provider=None, model=None, system=None:
                        next(v for k, v in replies.items() if k in prompt))
    df = gold_df()
    empty = {q: {"final": []} for q in df["qid"]}
    with tempfile.TemporaryDirectory() as d:
        llm_rows = run_eval_resumable(LLMOnly("ollama"), df, empty, Path(d) / "llm.csv",
                                      provider="ollama", label="llm_only", verbose=False)
    table = e1_table({"llm_only": llm_rows}).set_index("system")
    row = table.loc["llm_only"]
    assert row["abstention_accuracy"] == 0.0        # answered the unanswerable question
    assert row["false_refusal_rate"] == pytest.approx(1 / 3, abs=1e-3)   # refused Q3
    assert row["cites_gold_article"] == 0.5          # Q1 cites 5 (gold), Q2 cites 9 (not gold)
    assert pd.isna(row["citations_valid_first_attempt"])   # nothing to check against


def test_blind_sheet_hides_the_system_and_scores_are_mapped_back():
    from evaluation.manual_grading import write_grader_file
    from openpyxl import load_workbook

    df = gold_df()
    gold = pd.DataFrame({"question_id": df["qid"], "reference_answer": "ref", "supporting_text": "text"})

    def rows(system):
        return pd.DataFrame({"qid": df["qid"], "language": df["language"], "question": df["question"],
                             "expected": ["answer", "answer", "answer", "abstain"],
                             "answer": f"answer from {system}", "gold_source": df["gold_source"],
                             "gold_articles": ["5", "24", "7", ""], "final_articles": ""})

    table, key = blind_grading_table({"llm_only": rows("llm_only"), "naive_rag": rows("naive_rag")},
                                     gold, texts={}, seed=1)
    assert "system" not in table.columns and len(table) == 8 and table["item"].is_unique
    assert set(key["system"]) == {"llm_only", "naive_rag"}

    with tempfile.TemporaryDirectory() as d:
        paths = {}
        for grader, score in (("A", {"llm_only": 0, "naive_rag": 2}), ("B", {"llm_only": 1, "naive_rag": 2})):
            p = write_grader_file(table, Path(d) / f"g{grader}.xlsx", grader)
            wb = load_workbook(p)
            ws = wb["grading"]
            hdr = [c.value for c in ws[1]]
            sys_of = dict(zip(key["item"], key["system"]))
            for r in range(2, ws.max_row + 1):
                item = ws.cell(r, hdr.index("item") + 1).value
                ws.cell(r, hdr.index("correctness") + 1, score[sys_of[item]])
            wb.save(p)
            paths[grader] = p
        scores = manual_scores(paths["A"], paths["B"], key=key)
    assert scores["llm_only"]["manual_correctness_0_2"] == 0.5
    assert scores["naive_rag"]["manual_correctness_0_2"] == 2.0