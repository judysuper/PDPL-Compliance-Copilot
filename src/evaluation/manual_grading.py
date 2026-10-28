"""Manual grading: build blind grading sheets for two graders and measure their agreement.

Brief requirement: answer correctness on a 0/1/2 rubric, graded independently by two
team members, with agreement reported. We also grade whether the cited articles
support the claims (the automatic checker only proves they were retrieved).
"""

import json
import re
from pathlib import Path

import pandas as pd

CORRECTNESS = [0, 1, 2]
SUPPORT = ["yes", "partial", "no", "n/a"]
SOFT = ["yes", "no", "n/a"]

RUBRIC = [
    ("correctness = 2", "Answerable: correct and complete compared with the reference answer and the gold "
                        "article; no wrong legal claim. Unanswerable: the system clearly says the texts do "
                        "not answer it (NOT_COVERED or a plain statement), without invented specifics."),
    ("correctness = 1", "Answerable: main point right but an important condition, exception or number is "
                        "missing, or a minor inaccuracy. Unanswerable: hedges, but adds legal content that "
                        "reads like an answer."),
    ("correctness = 0", "Answerable: wrong, contradicts the article, or refuses although the article was "
                        "retrieved. Unanswerable: answers with invented or irrelevant specifics."),
    ("citation_support", "yes = every cited article supports the sentence it is attached to; partial = "
                         "some do; no = the cited articles do not support the claims; n/a = no citations "
                         "(abstained)."),
    ("soft_abstention", "Only for answers marked abstained=False: yes = the text says the information is "
                        "not in the sources but has no NOT_COVERED tag; no = it really answers; n/a = the "
                        "system abstained."),
    ("independence", "Grade alone. Do not look at the other grader's file or the automatic columns until "
                     "both files are finished."),
]

GRADER_COLS = ["correctness", "citation_support", "soft_abstention", "notes"]


def _norm_num(value):
    return re.sub(r"\.0+$", "", str(value).strip())


def load_article_texts(processed_dir):
    """{(document, lang, article_number): text} from data/processed/*_articles.csv (best effort)."""
    texts = {}
    for path in Path(processed_dir).glob("*_articles.csv"):
        parts = path.stem.split("_")            # pdpl_ar_articles -> pdpl, ar
        if len(parts) < 3:
            continue
        doc, lang = parts[0], parts[1]
        df = pd.read_csv(path, encoding="utf-8-sig", dtype=str, keep_default_na=False)
        num_col = next((c for c in df.columns if "article_number" in c), None)
        txt_col = next((c for c in df.columns if "article_text" in c), None)
        if not num_col or not txt_col:
            continue
        for num, txt in zip(df[num_col], df[txt_col]):
            texts[(doc, lang, _norm_num(num))] = txt
    return texts


def cited_texts(final_articles, cited_sources_order, language, texts, limit=1500):
    """Text of each source sent to the LLM, labelled like the prompt: [Source n] doc:article."""
    if not isinstance(final_articles, str) or not final_articles.strip():
        return ""
    blocks = []
    for n, ref in enumerate([r.strip() for r in final_articles.split("|")], start=1):
        doc, _, num = ref.partition(":")
        body = texts.get((doc, language, _norm_num(num))) or texts.get(
            (doc, "en" if language == "ar" else "ar", _norm_num(num)), "")
        blocks.append(f"[Source {n}] {doc} Article {num}\n{body[:limit]}")
    return "\n\n".join(blocks)


def build_grading_table(results, gold, article_texts=None):
    """One row per question: what a grader needs to see, nothing that reveals the metrics."""
    gold = gold.rename(columns={"question_id": "qid"})
    keep = [c for c in ["qid", "reference_answer", "supporting_text"] if c in gold.columns]
    df = results.merge(gold[keep], on="qid", how="left")
    texts = article_texts or {}
    df["sources_sent_to_llm"] = [
        cited_texts(fa, None, lang, texts) for fa, lang in zip(df.get("final_articles"), df["language"])
    ]
    table = pd.DataFrame({
        "qid": df["qid"],
        "language": df["language"],
        "question": df["question"],
        "expected": df["expected"].map({"answer": "answerable", "abstain": "not covered by the law"}),
        "system_answer": df["answer"],
        "reference_answer": df.get("reference_answer", ""),
        "gold_article": (df.get("gold_source", "").fillna("").astype(str) + " "
                         + df.get("gold_articles", "").fillna("").astype(str)
                         .map(lambda v: " | ".join(_norm_num(x) for x in v.split("|") if x.strip()))
                         ).str.strip(),
        "gold_article_text": df.get("supporting_text", "").fillna("").astype(str).str[:2000],
        "sources_sent_to_llm": df["sources_sent_to_llm"],
    })
    for col in GRADER_COLS:
        table[col] = ""
    return table.sort_values("qid").reset_index(drop=True)


def write_grader_file(table, path, grader):
    """Excel file with drop-downs, wrapped text and a rubric sheet."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        table.to_excel(xw, sheet_name="grading", index=False)
        pd.DataFrame(RUBRIC, columns=["field", "definition"]).to_excel(xw, sheet_name="rubric", index=False)
        ws, rub = xw.sheets["grading"], xw.sheets["rubric"]

        widths = {"qid": 8, "language": 8, "question": 45, "expected": 16, "system_answer": 70,
                  "reference_answer": 55, "gold_article": 16, "gold_article_text": 60,
                  "sources_sent_to_llm": 70, "correctness": 12, "citation_support": 15,
                  "soft_abstention": 15, "notes": 40}
        fill = PatternFill("solid", fgColor="FFF2CC")
        for idx, col in enumerate(table.columns, start=1):
            letter = ws.cell(row=1, column=idx).column_letter
            ws.column_dimensions[letter].width = widths.get(col, 20)
            ws.cell(row=1, column=idx).font = Font(bold=True)
            if col in GRADER_COLS:
                ws.cell(row=1, column=idx).fill = fill
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "D2"

        n = len(table) + 1
        for col, options in [("correctness", CORRECTNESS), ("citation_support", SUPPORT),
                             ("soft_abstention", SOFT)]:
            letter = ws.cell(row=1, column=list(table.columns).index(col) + 1).column_letter
            dv = DataValidation(type="list", formula1='"' + ",".join(map(str, options)) + '"',
                                allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{n}")
            for r in range(2, n + 1):
                ws[f"{letter}{r}"].fill = fill

        rub.column_dimensions["A"].width = 20
        rub.column_dimensions["B"].width = 110
        for row in rub.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        rub.cell(row=len(RUBRIC) + 3, column=1, value="grader")
        rub.cell(row=len(RUBRIC) + 3, column=2, value=grader)
    return path


def read_grades(path):
    df = pd.read_excel(path, sheet_name="grading", dtype={"qid": str, "item": str})
    df["correctness"] = pd.to_numeric(df["correctness"], errors="coerce")
    for col in ("citation_support", "soft_abstention"):
        df[col] = df[col].astype(str).str.strip().str.lower().replace({"nan": None, "": None})
    return df


def _kappa(a, b, labels, weights=None):
    """Cohen's kappa (weights=None or 'quadratic') without sklearn."""
    idx = {l: i for i, l in enumerate(labels)}
    k = len(labels)
    obs = [[0] * k for _ in range(k)]
    for x, y in zip(a, b):
        obs[idx[x]][idx[y]] += 1
    n = sum(map(sum, obs))
    if n == 0:
        return None
    row = [sum(r) for r in obs]
    col = [sum(obs[i][j] for i in range(k)) for j in range(k)]
    w = [[((i - j) ** 2 / (k - 1) ** 2 if weights == "quadratic" else (0 if i == j else 1))
          for j in range(k)] for i in range(k)]
    observed = sum(w[i][j] * obs[i][j] for i in range(k) for j in range(k)) / n
    expected = sum(w[i][j] * row[i] * col[j] for i in range(k) for j in range(k)) / n ** 2
    return None if expected == 0 else round(1 - observed / expected, 3)


def _gwet_ac1(a, b, labels):
    """Gwet's AC1: agreement corrected for chance, robust when one grade dominates.

    Cohen's kappa drops towards 0 when most items get the same grade (the "kappa
    paradox"); AC1 is the usual companion statistic in that case.
    """
    n = len(a)
    if n == 0:
        return None
    observed = sum(x == y for x, y in zip(a, b)) / n
    q = len(labels)
    pi = [(a.count(l) + b.count(l)) / (2 * n) for l in labels]
    expected = sum(p * (1 - p) for p in pi) / (q - 1)
    return None if expected == 1 else round((observed - expected) / (1 - expected), 3)


def agreement(grader_a, grader_b, key="qid"):
    """Agreement report + merged table with the disagreements to discuss.

    key: the row id shared by both files ("qid", or "item" for blind multi-system sheets).
    """
    a, b = read_grades(grader_a), read_grades(grader_b)
    left = list(dict.fromkeys([key, "qid", "language", "expected", "correctness",
                               "citation_support", "soft_abstention"]))
    m = a[[c for c in left if c in a.columns]].merge(
        b[[key, "correctness", "citation_support", "soft_abstention"]], on=key, suffixes=("_A", "_B"))

    both = m.dropna(subset=["correctness_A", "correctness_B"])
    ca, cb = both["correctness_A"].astype(int).tolist(), both["correctness_B"].astype(int).tolist()
    sup = m.dropna(subset=["citation_support_A", "citation_support_B"])
    sup = sup[(sup["citation_support_A"] != "n/a") & (sup["citation_support_B"] != "n/a")]

    report = {
        "n_graded_by_both": len(both),
        "n_missing": int(len(m) - len(both)),
        "correctness_exact_agreement": round(float((both["correctness_A"] == both["correctness_B"]).mean()), 3)
        if len(both) else None,
        "correctness_kappa_quadratic": _kappa(ca, cb, CORRECTNESS, "quadratic"),
        "correctness_kappa": _kappa(ca, cb, CORRECTNESS),
        "correctness_gwet_ac1": _gwet_ac1(ca, cb, CORRECTNESS),
        "a_lower_than_b": int(sum(x < y for x, y in zip(ca, cb))),
        "a_higher_than_b": int(sum(x > y for x, y in zip(ca, cb))),
        "mean_correctness_A": round(float(both["correctness_A"].mean()), 3) if len(both) else None,
        "mean_correctness_B": round(float(both["correctness_B"].mean()), 3) if len(both) else None,
        "mean_correctness_avg_of_2": round(float(((both["correctness_A"] + both["correctness_B"]) / 2).mean()), 3)
        if len(both) else None,
        "support_exact_agreement": round(float((sup["citation_support_A"] == sup["citation_support_B"]).mean()), 3)
        if len(sup) else None,
        "support_kappa": _kappa(sup["citation_support_A"].tolist(), sup["citation_support_B"].tolist(),
                                ["yes", "partial", "no"]) if len(sup) else None,
    }
    by_lang = {}
    for lang, g in both.groupby("language"):
        by_lang[lang] = round(float(((g["correctness_A"] + g["correctness_B"]) / 2).mean()), 3)
    report["mean_correctness_by_language"] = by_lang
    # compare as text so that "empty vs empty" is not counted as a disagreement
    def same(col):
        return m[f"{col}_A"].fillna("").astype(str) == m[f"{col}_B"].fillna("").astype(str)
    m["disagree"] = ~(same("correctness") & same("citation_support"))
    return report, m


def save_agreement(report, merged, out_dir):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "manual_grading_agreement.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    merged.to_csv(out_dir / "manual_grading_merged.csv", index=False, encoding="utf-8-sig")
    merged[merged["disagree"]].to_csv(out_dir / "manual_grading_disagreements.csv",
                                      index=False, encoding="utf-8-sig")
    return out_dir